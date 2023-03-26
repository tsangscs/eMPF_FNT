from datetime import datetime
import traceback
from openpyxl import load_workbook

class Func:
    functionID = ""
    dashboardModule = ""
    busProcess = ""
    timeBox = ""
    showAndTellStatus = ""
    numPassed = 0

    def __init__(self, functionID, dashboardModule, busProcess, timeBox, showAndTellStatus):
        self.functionID = functionID
        self.dashboardModule = dashboardModule
        self.busProcess = busProcess
        self.timeBox = timeBox
        self.showAndTellStatus = showAndTellStatus

    def print(self):
        print(str(self.functionID) + ": " + str(self.dashboardModule) + ", " +
              str(self.busProcess) + ", " + str(self.timeBox) + ", " +
              str(self.showAndTellStatus) + ", " + str(self.numPassed))


class ReleaseStatus:
    sheetName = "Release Status"
    colFunctionID = 1
    colDashboardModule = 3
    colBusProcess = 4
    colTimeBox = 6
    colNumPassed = 18
    colShowAndTell = 7
    funcDict = {}

    def loadFuncDict(self, workbook):
        try:
            sheet = workbook[self.sheetName]
            print(datetime.now().strftime("%H:%M:%S") + ": Processing [" + self.sheetName + "] row 2 to max row " + str(sheet.max_row))

            for rownum in range(2, sheet.max_row+1):
                funcID = sheet.cell(row=rownum, column=self.colFunctionID).value
                print("Loading row["+str(rownum)+"]: funcID = "+ funcID)
                if funcID is not None:
                    funcID = funcID.strip()
                else:
                    continue # Blank row
                dashboardModule = sheet.cell(row=rownum, column=self.colDashboardModule).value
                if dashboardModule is not None:
                    dashboardModule = dashboardModule.strip()
                else:
                    raise Exception("Blank module for " + funcID)
                busProcess = sheet.cell(row=rownum, column=self.colBusProcess).value
                if busProcess is not None:
                    busProcess = busProcess.strip()
                else:
                    raise Exception("Blank business process for " + funcID)
                timebox = sheet.cell(row=rownum, column=self.colTimeBox).value
                if timebox is not None:
                    timebox = timebox.strip()
                else:
                    raise Exception("Blank timebox for " + funcID)
                showAndTell = sheet.cell(row=rownum, column=self.colShowAndTell).value
                if showAndTell is not None:
                    if showAndTell == "":
                        showAndTell = "Not Completed"
                    else:
                        showAndTell = showAndTell.strip()
                else:
                    showAndTell = "Not Completed"

                func = Func(funcID, dashboardModule, busProcess, timebox, showAndTell)
                self.funcDict[funcID] = func

        except Exception as ex:
            print(ex)
            traceback.print_exc()

    def printFuncDict(self):
        for f in self.funcDict.values():
            f.print()

    def updateReleaseStatusSheet(self, workbook):
        try:
            sheet = workbook[self.sheetName]
            print(datetime.now().strftime("%H:%M:%S") + ": Processing [" + self.sheetName + "] row 2 to max row " + str(sheet.max_row))

            for rownum in range(2, sheet.max_row+1):
                funcID = sheet.cell(row=rownum, column=self.colFunctionID).value
                if funcID is not None:
                    funcID = funcID.strip()
                    func = self.funcDict[funcID]
                    sheet.cell(row=rownum, column=self.colNumPassed).value = func.numPassed
                else:
                    continue

        except Exception as ex:
            print(ex)
            traceback.print_exc()

class TestCase:
    testName = ""
    dashboardModule = ""
    busProcess = ""
    timeBox = ""
    passStatus = ""
    showAndTellStatus = ""
    funclist = {}
    error = ""

    def __init__(self, testName):
        self.testName = testName

    def print(self):
        print(self.testName + ", " +
              self.dashboardModule + ", " +
              self.busProcess + ", " +
              self.timeBox + ", " +
              self.passStatus + ", " +
              self.showAndTellStatus + ", " +
              ":".join(self.funclist)
              )

    def appendError(self, errMsg):
        if self.error != "":
            self.error += ": "
        self.error += errMsg

class TestCaseSheet:
    sheetName = "Main Sheet"
    colTestName = 1
    colFNTStatus = 3
    colModule = 4
    colBusProcess = 6
    colRQID = 7
    colShowAndTellStatus = 8
    colTimeBox = 21
    colErrMsg = 22

    def updateTestCaseSheet(self, workbook):
        try:
            sheet = workbook[self.sheetName]
            sheet.cell(row=1, column=self.colTimeBox).value = "Time Box"
            sheet.cell(row=1, column=self.colErrMsg).value = "Error Message"

            #Read Function Release Status
            releaseStatus = ReleaseStatus()
            releaseStatus.loadFuncDict(workbook)
            print(datetime.now().strftime("%H:%M:%S") + ": Processing [" + self.sheetName + "] row 2 to max row " + str(sheet.max_row))

            for rownum in range(2, sheet.max_row+1):
                testName = sheet.cell(row=rownum, column=self.colTestName).value
                if testName is None:
                    continue
                elif len(testName.strip()) == 0:
                    continue
                try:
                    testCase = TestCase(testName)

                    # Check Pass Status
                    testCase.passStatus = sheet.cell(row=rownum, column=self.colFNTStatus).value

                    # Check Functions
                    rqidlist = sheet.cell(row=rownum, column=self.colRQID).value
                    if rqidlist is not None:
                        rqidlist = rqidlist.strip()
                        if len(rqidlist) != 0 and rqidlist != "#N/A":
    #                        if rqidlist[len(rqidlist)] == ",":
    #                            rqidlist = rqidlist[0:len(rqidlist)-2]
                            rqidlist = rqidlist.replace(chr(10), ",")
                            testCase.funclist = rqidlist.split(",")

                    if testCase.funclist is None or len(testCase.funclist) == 0:
                        testCase.appendError("RQID is #N/A or blank")
                    else:
                        dashboardModule = []
                        busProcess = []
                        for rqid in testCase.funclist:
                            funcID = rqid.strip()
                            if len(funcID) == 0:
                                continue
                            func = releaseStatus.funcDict.get(funcID)
                            # Found Functions
                            if func is not None:
                                # Update Dashboard Module
                                try:
                                    dashboardModule.index(func.dashboardModule)
                                except Exception as ex:
                                    dashboardModule.append(func.dashboardModule)
                                    dashboardModule.sort()
                                    testCase.dashboardModule = ":".join(dashboardModule)

                                # Update Business Process
                                try:
                                    busProcess.index(func.busProcess)
                                except Exception as ex:
                                    busProcess.append(func.busProcess)
                                    busProcess.sort()
                                    testCase.busProcess = ":".join(busProcess)

                                # Update Test Case Timebox
                                if func.timeBox in ["Day 2", "De-scoped"]:
                                    testCase.timeBox = func.timeBox
                                elif func.timeBox == "TX13":
                                    if testCase.timeBox in ["", "TX1-10", "TX-11", "TX-12"]:
                                        testCase.timeBox = func.timeBox
                                elif func.timeBox == "TX12":
                                    if testCase.timeBox in ["", "TX1-10", "TX-11"]:
                                        testCase.timeBox = func.timeBox
                                elif func.timeBox == "TX11":
                                    if testCase.timeBox in ["", "TX1-10"]:
                                        testCase.timeBox = func.timeBox
                                elif func.timeBox == "TX1-10":
                                    if testCase.timeBox == "":
                                        testCase.timeBox = func.timeBox
                                else:
                                    testCase.appendError("Invalid timebox of " + funcID + " = " + func.timeBox)

                                # Update Test Case Show & Tell Status
                                if func.showAndTellStatus == "Not Completed":
                                    testCase.showAndTellStatus = "Not Completed"
                                elif func.showAndTellStatus in ["N/A", "Done"]:
                                    if testCase.showAndTellStatus != "Not Completed":
                                        testCase.showAndTellStatus = "Done"
                                else:
                                    testCase.appendError("Invalid S&T Status of " + funcID + " = " + func.showAndTellStatus)

                                # Update function pass status
                                if testCase.passStatus[0:6] == "Passed":
                                    func.numPassed += 1
                                    releaseStatus.funcDict.update({funcID: func})
                            else:
                                testCase.appendError(funcID + " not valid")

                        # Update test case
                        sheet.cell(row=rownum, column=self.colModule).value = testCase.dashboardModule
                        sheet.cell(row=rownum, column=self.colBusProcess).value = testCase.busProcess
                        sheet.cell(row=rownum, column=self.colTimeBox).value = testCase.timeBox
                        sheet.cell(row=rownum, column=self.colShowAndTellStatus).value = testCase.showAndTellStatus
                    # end if
                    sheet.cell(row=rownum, column=self.colErrMsg).value = testCase.error

                except Exception as ex:
                    errMsg = str(ex)
                    print(datetime.now().strftime("%H:%M:%S") +
                          ": Error found in [" + self.sheetName + "] at row " + str(rownum) + " Error: " + errMsg)

            # end for

            #print("Result Function Status: ")
            #releaseStatus.printFuncDict()
            releaseStatus.updateReleaseStatusSheet(workbook)

        except Exception as ex:
            errMsg = str(ex)
            print("Error found at row[" + str(rownum) + "]: " + errMsg)
            traceback.print_exc()
            if testCase is not None:
                if testCase.error is not None:
                    errMsg = errMsg + ", " + testCase.error
            sheet.cell(row=rownum, column=self.colErrMsg).value = testCase.error

# Press the green button in the gutter to run the script.
if __name__ == '__main__':
    infile = input("Enter filename of FNT Execution Report: ")
    extensionIdx = infile.rfind(".xlsx")
    outfile = infile[0:extensionIdx] + "_out.xlsx"
    print("Output file = " + outfile)

    print(datetime.now().strftime("%H:%M:%S") + ": Loading input workbook ...")
    workbook = load_workbook(filename=infile)

    print(datetime.now().strftime("%H:%M:%S") + ": Processing ...")
    testCaseSheet = TestCaseSheet()
    testCaseSheet.updateTestCaseSheet(workbook)

    print(datetime.now().strftime("%H:%M:%S") + ": Saving output workbook ...")
    workbook.save(outfile)
    print(datetime.now().strftime("%H:%M:%S") + ": Job completed!")
